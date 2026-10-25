from __future__ import annotations

import argparse
from copy import copy
import datetime as dt
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook


def json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def xml(value):
    return ET.tostring(value.to_tree(), encoding="unicode") if value is not None else None


def expected_argument(value):
    address, serialized = value.split("=", 1)
    sheet, cell = address.rsplit("!", 1)
    return (sheet, cell.upper()), json.loads(serialized)


def recalculate(saved_path, expected):
    executable = shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice is required for --recalculate")
    with tempfile.TemporaryDirectory(prefix="genoffice-xlsx-verify-") as temp:
        temp = Path(temp)
        result = subprocess.run([
            executable,
            f"-env:UserInstallation={(temp / 'profile').as_uri()}",
            "--headless", "--convert-to", "xlsx", "--outdir", str(temp), str(saved_path.resolve()),
        ], capture_output=True, text=True, timeout=60)
        output = temp / saved_path.name
        if result.returncode or not output.is_file():
            raise RuntimeError(f"LibreOffice calculation failed: {result.stdout} {result.stderr}")
        computed = load_workbook(output, data_only=True)
        formulas = load_workbook(saved_path, data_only=False)
        errors = []
        values = {}
        try:
            for sheet in formulas:
                for cell in sheet._cells.values():
                    if cell.data_type != "f":
                        continue
                    value = computed[sheet.title][cell.coordinate]
                    qualified = f"{sheet.title}!{cell.coordinate}"
                    values[qualified] = json_value(value.value)
                    if value.value is None or value.data_type == "e":
                        errors.append({"kind": "recalculated_formula_error", "cell": qualified, "value": json_value(value.value)})
            for (sheet, cell), wanted in expected.items():
                actual = computed[sheet][cell].value
                matches = abs(actual - wanted) < 1e-8 if isinstance(actual, (int, float)) and isinstance(wanted, (int, float)) else json_value(actual) == wanted
                if not matches:
                    errors.append({"kind": "recalculated_value_mismatch", "cell": f"{sheet}!{cell}", "expected": wanted, "actual": json_value(actual)})
            return {"engine": "LibreOffice", "status": "passed" if not errors else "failed", "formula_count": len(values), "values": values, "errors": errors, "method": "Converted a temporary copy with an isolated LibreOffice profile; saved plugin output was not rewritten."}
        finally:
            computed.close()
            formulas.close()


def verify(original_path, saved_path, expected, allowed_entries, require_cache):
    before_hash = hashlib.sha256(original_path.read_bytes()).hexdigest()
    errors = []
    original = load_workbook(original_path, data_only=False)
    saved = load_workbook(saved_path, data_only=False)
    cached = load_workbook(saved_path, data_only=True)
    inspected_cells = 0
    formulas = []
    features = {}
    try:
        if original.sheetnames != saved.sheetnames:
            errors.append({"kind": "sheet_names", "before": original.sheetnames, "after": saved.sheetnames})
        for left in original:
            if left.title not in saved:
                continue
            right = saved[left.title]
            feature_checks = {
                "sheet_state": (left.sheet_state, right.sheet_state),
                "merged_cells": (sorted(str(value) for value in left.merged_cells.ranges), sorted(str(value) for value in right.merged_cells.ranges)),
                "freeze_panes": (left.freeze_panes, right.freeze_panes),
                "protection": (xml(left.protection), xml(right.protection)),
                "data_validations": (xml(left.data_validations), xml(right.data_validations)),
                "conditional_formats": ([xml(rule) for rules in left.conditional_formatting._cf_rules.values() for rule in rules], [xml(rule) for rules in right.conditional_formatting._cf_rules.values() for rule in rules]),
                "chart_count": (len(left._charts), len(right._charts)),
                "image_count": (len(left._images), len(right._images)),
                "table_names": (sorted(left.tables), sorted(right.tables)),
            }
            features[left.title] = {key: a == b for key, (a, b) in feature_checks.items()}
            for key, (a, b) in feature_checks.items():
                if a != b:
                    errors.append({"kind": key, "sheet": left.title, "before": a, "after": b})
            for coordinate in sorted(set(left._cells) | set(right._cells)):
                a = left.cell(*coordinate)
                b = right.cell(*coordinate)
                wanted = expected.get((left.title, a.coordinate), json_value(a.value))
                inspected_cells += 1
                if json_value(b.value) != wanted:
                    errors.append({"kind": "cell_value", "cell": f"{left.title}!{a.coordinate}", "expected": wanted, "actual": json_value(b.value)})
                for field in ["font", "fill", "border", "alignment", "protection", "number_format"]:
                    if copy(getattr(a, field)) != copy(getattr(b, field)):
                        errors.append({"kind": "cell_style", "cell": f"{left.title}!{a.coordinate}", "field": field})
                if (a.comment.text if a.comment else None) != (b.comment.text if b.comment else None):
                    errors.append({"kind": "cell_comment", "cell": f"{left.title}!{a.coordinate}"})
                if (a.hyperlink.target if a.hyperlink else None) != (b.hyperlink.target if b.hyperlink else None):
                    errors.append({"kind": "cell_hyperlink", "cell": f"{left.title}!{a.coordinate}"})
                if b.data_type == "f":
                    value = cached[left.title][b.coordinate]
                    formulas.append({"cell": f"{left.title}!{b.coordinate}", "formula": b.value, "cached": json_value(value.value)})
                    if value.data_type == "e":
                        errors.append({"kind": "formula_error", "cell": f"{left.title}!{b.coordinate}", "value": value.value})
                    if require_cache and value.value is None:
                        errors.append({"kind": "missing_formula_cache", "cell": f"{left.title}!{b.coordinate}"})
            for dimension_name in ["row_dimensions", "column_dimensions"]:
                a = getattr(left, dimension_name)
                b = getattr(right, dimension_name)
                for key in set(a) | set(b):
                    if key not in a or key not in b or dict(a[key]) != dict(b[key]):
                        errors.append({"kind": dimension_name, "sheet": left.title, "key": str(key)})
        for (sheet, cell), value in expected.items():
            if sheet not in saved or json_value(saved[sheet][cell].value) != value:
                errors.append({"kind": "expected_edit_missing", "cell": f"{sheet}!{cell}", "expected": value})
    finally:
        original.close()
        saved.close()
        cached.close()
    with ZipFile(original_path) as a, ZipFile(saved_path) as b:
        left_names = {entry.filename for entry in a.infolist() if not entry.is_dir()}
        right_names = {entry.filename for entry in b.infolist() if not entry.is_dir()}
        changed = sorted(name for name in left_names & right_names if a.read(name) != b.read(name))
        identical = sorted(name for name in left_names & right_names if a.read(name) == b.read(name))
        missing = sorted(left_names - right_names)
        added = sorted(right_names - left_names)
        permitted = set(allowed_entries) | {"xl/workbook.xml", "xl/calcChain.xml", "docProps/core.xml"}
        for name in changed + missing:
            if not name.startswith("xl/worksheets/sheet") and name not in permitted:
                errors.append({"kind": "unrelated_ooxml_changed_or_missing", "entry": name})
        for name in [entry for entry in left_names if entry.startswith(("xl/charts/", "xl/media/", "xl/drawings/", "customXml/"))]:
            if name not in right_names or a.read(name) != b.read(name):
                errors.append({"kind": "embedded_part_not_preserved", "entry": name})
    after_hash = hashlib.sha256(original_path.read_bytes()).hexdigest()
    if before_hash != after_hash:
        errors.append({"kind": "input_was_modified"})
    return {
        "status": "passed" if not errors else "failed",
        "original": str(original_path),
        "saved": str(saved_path),
        "original_sha256": before_hash,
        "saved_sha256": hashlib.sha256(saved_path.read_bytes()).hexdigest(),
        "input_unchanged": before_hash == after_hash,
        "inspected_cells": inspected_cells,
        "expected_edits": {f"{sheet}!{cell}": value for (sheet, cell), value in expected.items()},
        "features": features,
        "formula_count": len(formulas),
        "formulas": formulas,
        "ooxml": {"identical_count": len(identical), "changed": changed, "missing": missing, "added": added, "identical": identical},
        "errors": errors,
    }


def main():
    parser = argparse.ArgumentParser(description="Read-only independent XLSX round-trip validation")
    parser.add_argument("original", type=Path)
    parser.add_argument("saved", type=Path)
    parser.add_argument("--expect", action="append", default=[], help='Sheet!A1=JSON value (repeatable)')
    parser.add_argument("--allow-entry", action="append", default=[])
    parser.add_argument("--require-cached", action="store_true")
    parser.add_argument("--recalculate", action="store_true")
    parser.add_argument("--expect-calculated", action="append", default=[], help='Sheet!A1=JSON calculated value (repeatable)')
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = verify(args.original, args.saved, dict(expected_argument(value) for value in args.expect), args.allow_entry, args.require_cached)
    if args.recalculate:
        report["recalculation"] = recalculate(args.saved, dict(expected_argument(value) for value in args.expect_calculated))
        report["errors"].extend(report["recalculation"]["errors"])
        report["status"] = "passed" if not report["errors"] else "failed"
    serialized = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(serialized)
    print(serialized)
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    sys.exit(main())
