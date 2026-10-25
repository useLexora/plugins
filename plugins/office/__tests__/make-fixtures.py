from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image as PillowImage


ROOT = Path(__file__).resolve().parent.parent
FIXTURES = ROOT / "fixtures"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SYNTHETIC = "所有名称、数量、单价及日期均为插件验收合成数据，不代表真实业务。"
EXPECTED = {
    "销售明细!E4": 258,
    "销售明细!E5": 247.5,
    "销售明细!E6": 225,
    "销售明细!E7": 112,
    "销售明细!E8": 299,
    "销售明细!E9": 119.4,
    "销售明细!C10": 21,
    "销售明细!E10": 1260.9,
    "汇总!B4": 1260.9,
    "汇总!B5": 126.09,
    "汇总!B6": 1134.81,
    "汇总!B7": 21,
}


def style(sheet):
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=11, color="172B4D")
            cell.alignment = Alignment(vertical="center")
    sheet.sheet_view.showGridLines = False


def business():
    book = Workbook()
    sheet = book.active
    sheet.title = "销售明细"
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "秋季文具采购 · 合成测试工作簿"
    sheet.merge_cells("A2:H2")
    sheet["A2"] = SYNTHETIC + " 蓝色文字为输入；金额和汇总为公式。"
    headers = ["编号", "商品", "数量", "单价（元）", "金额（元）", "日期", "状态", "备注"]
    sheet.append(headers)
    rows = [
        ["SYN-001", "晨光钢笔", 2, 129, "已确认", "文字与数字编辑"],
        ["SYN-002", "方格笔记本", 5, 49.5, "待发货", "Unicode：中文 / café / 😀"],
        ["SYN-003", "桌面收纳盒", 3, 75, "已确认", "含跨表汇总"],
        ["SYN-004", "便签纸", 4, 28, "待发货", "合成测试数据"],
        ["SYN-005", "阅读台灯", 1, 299, "已确认", "合成测试数据"],
        ["SYN-006", "彩色胶带", 6, 19.9, "已确认", "合成测试数据"],
    ]
    for r, item in enumerate(rows, 4):
        code, name, count, price, status, note = item
        sheet.append([code, name, count, price, f"=C{r}*D{r}", dt.datetime(2026, 10, r + 10), status, note])
    sheet["B10"] = "合计"
    sheet["C10"] = "=SUM(C4:C9)"
    sheet["E10"] = "=SUM(E4:E9)"
    sheet["B4"].comment = Comment("合成商品名称。可将本格改为“晨光钢笔（已修改）”测试文本写回。", "Fixture")
    sheet["H4"].hyperlink = "https://example.com/synthetic-order/001"
    sheet.freeze_panes = "C4"
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 34
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 24
    for col in ["C", "D", "E", "F", "G"]:
        sheet.column_dimensions[col].width = 17
    sheet.column_dimensions["H"].width = 42
    style(sheet)
    sheet["A1"].font = Font(name="Arial", size=20, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="16665C")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    for cell in sheet[3]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234D64")
    for row in sheet.iter_rows(min_row=4, max_row=9):
        for cell in row:
            cell.border = Border(bottom=Side(style="hair", color="D8E6E4"))
            cell.font = Font(name="Arial", size=11, color="000000" if cell.data_type == "f" else "0000FF")
        row[2].fill = PatternFill("solid", fgColor="FFF2CC")
        row[3].fill = PatternFill("solid", fgColor="FFF2CC")
        row[3].number_format = '¥#,##0.00;[Red](¥#,##0.00)'
        row[4].number_format = '¥#,##0.00;[Red](¥#,##0.00)'
        row[5].number_format = "yyyy-mm-dd"
    sheet["E10"].number_format = '¥#,##0.00'
    table = Table(displayName="SyntheticSales", ref="A3:H9")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    quantity = DataValidation(type="whole", operator="between", formula1="0", formula2="10000")
    quantity.errorTitle = "数量无效"
    quantity.error = "请输入 0 到 10000 的整数。"
    quantity.showErrorMessage = True
    sheet.add_data_validation(quantity)
    quantity.add("C4:C9")
    status = DataValidation(type="list", formula1='"已确认,待发货"')
    sheet.add_data_validation(status)
    status.add("G4:G9")
    sheet.conditional_formatting.add("E4:E9", CellIsRule(operator="greaterThan", formula=["200"], fill=PatternFill("solid", fgColor="D7EFE2")))
    chart = BarChart()
    chart.title = "合成商品采购金额"
    chart.y_axis.title = "元"
    chart.add_data(Reference(sheet, min_col=5, min_row=3, max_row=9), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=2, min_row=4, max_row=9))
    chart.height = 8
    chart.width = 15
    sheet.add_chart(chart, "J6")
    image_bytes = io.BytesIO()
    PillowImage.new("RGB", (200, 48), "#16665c").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    sheet.add_image(Image(image_bytes), "J2")
    summary = book.create_sheet("汇总")
    summary.append(["采购汇总", "金额 / 数量"])
    summary.append(["说明", SYNTHETIC])
    summary.append(["公式说明", "按明细合计与参数页折扣率计算"])
    summary.append(["原始总额", "='销售明细'!E10"])
    summary.append(["折扣额", "=B4*'参数'!B2"])
    summary.append(["折后总额", "=B4-B5"])
    summary.append(["总数量", "='销售明细'!C10"])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 70
    style(summary)
    for r in range(4, 7):
        summary[f"B{r}"].number_format = '¥#,##0.00'
        summary[f"B{r}"].font = Font(name="Arial", size=11, color="008000")
    params = book.create_sheet("参数")
    params.append(["参数", "值", "来源"])
    params.append(["合成折扣率", 0.1, SYNTHETIC])
    style(params)
    params["B2"].number_format = "0.0%"
    params.sheet_state = "hidden"
    book.save(FIXTURES / "business.xlsx")


def edge_cases():
    book = Workbook()
    book.active.title = "空白"
    book.save(FIXTURES / "empty.xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "稀疏数据"
    sheet["A1"] = SYNTHETIC
    sheet["B100000"] = "高行稀疏文本"
    sheet["XFD1048576"] = 42
    for cell in [sheet["A1"], sheet["B100000"], sheet["XFD1048576"]]:
        cell.font = Font(name="Arial", size=11)
    book.save(FIXTURES / "sparse.xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "受保护"
    sheet["A1"] = SYNTHETIC
    sheet["A2"] = "锁定内容"
    sheet["B2"] = "解锁输入"
    sheet["B2"].protection = Protection(locked=False)
    sheet.protection.sheet = True
    sheet.protection.set_password("synthetic-fixture")
    style(sheet)
    book.save(FIXTURES / "protected.xlsx")
    (FIXTURES / "invalid.xlsx").write_bytes("这不是 ZIP，也不是有效 XLSX。仅用于错误分支测试。".encode())


def calculate_caches(path):
    executable = shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice is required to calculate fixture formulas")
    with tempfile.TemporaryDirectory(prefix="genoffice-xlsx-fixture-") as temp:
        temp = Path(temp)
        result = subprocess.run([
            executable,
            f"-env:UserInstallation={(temp / 'profile').as_uri()}",
            "--headless", "--convert-to", "xlsx", "--outdir", str(temp), str(path),
        ], capture_output=True, text=True, timeout=60)
        calculated = temp / path.name
        if result.returncode or not calculated.is_file():
            raise RuntimeError(f"LibreOffice calculation failed: {result.stdout} {result.stderr}")
        values = load_workbook(calculated, data_only=True)
        for qualified, expected in EXPECTED.items():
            sheet, address = qualified.split("!")
            actual = values[sheet][address].value
            if not isinstance(actual, (int, float)) or abs(actual - expected) > 1e-8:
                raise AssertionError(f"Unexpected calculated value {qualified}: {actual} != {expected}")
        values.close()
        with ZipFile(path) as original, ZipFile(calculated) as computed:
            output = io.BytesIO()
            with ZipFile(output, "w", ZIP_DEFLATED) as target:
                for entry in original.infolist():
                    content = original.read(entry.filename)
                    if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", entry.filename):
                        cells = {}
                        for cell in ET.fromstring(computed.read(entry.filename)).findall(".//s:c", NS):
                            value = cell.find("s:v", NS)
                            if cell.find("s:f", NS) is not None and value is not None:
                                cells[cell.attrib["r"]] = value.text
                        def replace(match):
                            block = match.group(0)
                            address = re.search(r'\br="([A-Z]+[0-9]+)"', block).group(1)
                            if "<f" not in block or address not in cells:
                                return block
                            value = f"<v>{cells[address]}</v>"
                            return re.sub(r"<v(?:\s[^>]*)?>.*?</v>|<v\s*/>", value, block, count=1, flags=re.S)
                        content = re.sub(r"<c\b[^>]*?(?<!/)>.*?</c>", replace, content.decode(), flags=re.S).encode()
                    target.writestr(entry, content)
                target.writestr("customXml/fixture-preservation.xml", '<fixture xmlns="urn:lexora:xlsx:test">保留未知部件：合成测试数据</fixture>')
            path.write_bytes(output.getvalue())
    cached = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    try:
        formula_count = 0
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        formula_count += 1
                        value = cached[sheet.title][cell.coordinate]
                        if value.value is None or value.data_type == "e":
                            raise AssertionError(f"Missing or erroneous formula cache: {sheet.title}!{cell.coordinate}")
        return {"engine": "LibreOffice", "formula_count": formula_count, "formula_errors": 0, "expected_values": EXPECTED, "cache_transfer": "Only formula v elements copied into original OOXML; styles, charts, drawings and formulas preserved."}
    finally:
        cached.close()
        formulas.close()


def main():
    FIXTURES.mkdir(exist_ok=True)
    business()
    edge_cases()
    calculation = calculate_caches(FIXTURES / "business.xlsx")
    manifest = {
        "synthetic": True,
        "description": SYNTHETIC,
        "calculation": calculation,
        "files": {file.name: {"bytes": file.stat().st_size, "sha256": hashlib.sha256(file.read_bytes()).hexdigest()} for file in sorted(FIXTURES.glob("*.xlsx"))},
        "recommended_edits": {"business.xlsx": {"销售明细!B4": "晨光钢笔（已修改）", "销售明细!C4": 3}},
        "expected_after_edits": {"销售明细!E4": 387, "销售明细!C10": 22, "销售明细!E10": 1389.9, "汇总!B4": 1389.9, "汇总!B5": 138.99, "汇总!B6": 1250.91, "汇总!B7": 22},
    }
    (FIXTURES / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
